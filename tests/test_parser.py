from io import BytesIO

import pytest
from docx import Document

from xassemble.parser import QuestionnaireError, parse_questionnaire

from .helpers import questionnaire_docx


def test_parses_questionnaire_tables() -> None:
    schema = parse_questionnaire(questionnaire_docx())
    assert schema.variables == {"client_name", "urgent", "reason"}
    assert schema.questions[2].skip_if == "urgent == false"
    assert schema.outputs[0].label == "letter"
    assert schema.questions[0].commentary == ""


def test_question_text_paragraphs_after_the_first_become_commentary() -> None:
    content = questionnaire_docx()
    document = Document(BytesIO(content))
    table = document.tables[0]
    cell = table.rows[1].cells[1]
    cell.text = "Client name"
    cell.add_paragraph("Use the full legal name.")
    cell.add_paragraph("Include any trading name too.")
    destination = BytesIO()
    document.save(destination)

    schema = parse_questionnaire(destination.getvalue())
    question = next(q for q in schema.questions if q.variable_name == "client_name")
    assert question.question_text == "Client name"
    assert question.commentary == "Use the full legal name.\nInclude any trading name too."


def test_choice_must_have_options() -> None:
    content = questionnaire_docx(
        questions=[["kind", "Which kind?", "choice", "", "", "", ""]]
    )
    with pytest.raises(QuestionnaireError, match="choice requires options"):
        parse_questionnaire(content)


def test_skip_condition_can_only_reference_earlier_questions() -> None:
    content = questionnaire_docx(
        questions=[
            ["first", "First", "text", "", "", "later == 'x'", ""],
            ["later", "Later", "text", "", "", "", ""],
        ]
    )
    with pytest.raises(QuestionnaireError, match="Unknown variable.*later"):
        parse_questionnaire(content)


def test_duplicate_variables_are_rejected() -> None:
    content = questionnaire_docx(
        questions=[
            ["name", "Name", "text", "", "", "", ""],
            ["name", "Name again", "text", "", "", "", ""],
        ]
    )
    with pytest.raises(QuestionnaireError, match="Duplicate question variable"):
        parse_questionnaire(content)

