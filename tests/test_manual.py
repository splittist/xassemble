from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from xassemble.app import DOCX_MEDIA_TYPE
from xassemble.manual import ManualExportError, export_manual_template
from xassemble.parser import parse_questionnaire

from .auth_helpers import authenticated_app, log_in
from .helpers import questionnaire_docx, template_docx


def schema():
    return parse_questionnaire(questionnaire_docx(questions=[
        ["client_name", "What is the name of the client?\nUse the registered name.",
         "text", "", "", "", ""],
        ["urgent", "Is this urgent?", "yesno", "", "", "", ""],
    ]))


def save(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_split_runs_preserve_formatting_tabs_headers_tables_and_package_parts():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Dear ").italic = True
    paragraph.add_run("{{ cli").bold = True
    paragraph.add_run("ent_name }} and welcome.")
    paragraph.add_run("\tAfter tab\nAfter break")
    document.sections[0].header.paragraphs[0].text = "{{ project_code }}"
    document.add_table(rows=1, cols=1).cell(0, 0).text = "{{ client_name | upper }}"
    original = save(document)
    result = export_manual_template(original, schema())
    rendered = Document(BytesIO(result))
    paragraph = rendered.paragraphs[0]
    assert paragraph.text == "Dear [the name of the client] and welcome.\tAfter tab\nAfter break"
    placeholder = next(run for run in paragraph.runs if run.text.startswith("[the"))
    assert placeholder.bold
    assert placeholder.font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert paragraph.runs[0].italic
    assert rendered.sections[0].header.paragraphs[0].text == "[project code]"
    assert rendered.tables[0].cell(0, 0).text == "[the name of the client (uppercase)]"
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(result)) as after:
        assert before.namelist() == after.namelist()
        assert before.read("word/styles.xml") == after.read("word/styles.xml")


def test_nested_branches_and_alternatives_keep_all_text_with_correct_shading():
    document = Document()
    for text in ["{%p if urgent %}", "Urgent {{ client_name }}",
                 "{% if client_name %}Named{% else %}Unnamed{% endif %}",
                 "{%p elif client_name == 'Other' %}", "Other client",
                 "{%p else %}", "Normal", "{%p endif %}", "Always"]:
        document.add_paragraph(text)
    rendered = Document(BytesIO(export_manual_template(save(document), schema(),
                                                       include_if="urgent == true")))
    text = "\n".join(p.text for p in rendered.paragraphs)
    for expected in ["Urgent [the name of the client]", "Named", "Unnamed", "Other client",
                     "Normal", "Always", "only within C1", "none of C1, C4 applies",
                     "Use this whole document only if “Is this urgent” is Yes"]:
        assert expected in text
    assert "{%" not in text and "{{" not in text
    urgent = rendered.paragraphs[1].runs[0]
    assert urgent._r.rPr.find(qn("w:shd")).get(qn("w:fill")) == "DDEBF7"
    assert rendered.paragraphs[8].runs[0]._r.rPr is None
    named = next(r for r in rendered.paragraphs[2].runs if r.text == "Named")
    assert named._r.rPr.find(qn("w:shd")).get(qn("w:fill")) == "E2EFDA"


@pytest.mark.parametrize("text, message", [
    ("{% for item in items %}x{% endfor %}", "Unsupported manual-export tag"),
    ("{% if urgent %}x", "Unclosed conditional"),
    ("{% endif %}", "Unexpected endif"),
    ("{{ missing }}", "No questionnaire question"),
    ("{{ client_name", "Incomplete template tag"),
    ("{% if urgent %}x{% else %}y{% else %}z{% endif %}", "Unexpected else"),
])
def test_unsupported_or_broken_templates_fail_explicitly(text, message):
    with pytest.raises(ManualExportError, match=message):
        export_manual_template(template_docx(text), schema())


def test_whitespace_controls_comments_and_calculated_fields():
    result = export_manual_template(template_docx(
        "{# internal #}{%- if urgent -%}{{ client_name | replace('a', 'b') }}"
        "{%- endif -%}"), schema())
    text = Document(BytesIO(result)).paragraphs[0].text
    assert "internal" not in text
    assert "[Calculate / review manually: client_name | replace('a', 'b')]" in text


def test_arithmetic_signs_are_not_confused_with_whitespace_control():
    result = export_manual_template(template_docx("{{ -5 }} {{- -5 -}}"), schema())
    text = Document(BytesIO(result)).paragraphs[0].text
    assert text == "[Calculate / review manually: -5] [Calculate / review manually: -5]"


def test_manual_download_uses_current_published_versions_and_requires_login(tmp_path):
    app, _ = authenticated_app(tmp_path / "manual.sqlite3")
    with TestClient(app) as client:
        path = "/document-sets/letters/current/templates/letter/manual"
        assert client.get(path).status_code == 401
        log_in(client)
        client.post("/document-sets", json={"slug": "letters", "name": "Letters"})
        assert client.get(path).status_code == 404
        uploaded = client.post("/document-sets/letters/questionnaire-versions",
                               files={"file": ("q.docx", questionnaire_docx(), DOCX_MEDIA_TYPE)})
        client.post(f"/document-sets/letters/questionnaire-versions/{uploaded.json()['id']}/publish")
        assert client.get(path).status_code == 404
        uploaded = client.post("/document-sets/letters/template-versions", data={"label": "letter"},
                               files={"file": ("t.docx", template_docx(), DOCX_MEDIA_TYPE)})
        client.post(f"/document-sets/letters/template-versions/{uploaded.json()['id']}/publish")
        # An unpublished replacement must not affect the exported copy.
        client.post("/document-sets/letters/template-versions", data={"label": "letter"},
                    files={"file": ("t.docx", template_docx("Draft"), DOCX_MEDIA_TYPE)})
        response = client.get(path)
        assert response.status_code == 200
        assert 'letter-manual-t1-q1.docx' in response.headers['content-disposition']
        assert Document(BytesIO(response.content)).paragraphs[0].text == "Dear [Client name]"
        # Source download remains the original assembly template.
        source = client.get(path.removesuffix("/manual"))
        assert Document(BytesIO(source.content)).paragraphs[0].text == "Dear {{ client_name }}"
