from __future__ import annotations

from io import BytesIO

from docx import Document


def questionnaire_docx(
    questions: list[list[str]] | None = None,
    outputs: list[list[str]] | None = None,
) -> bytes:
    document = Document()
    question_headers = [
        "variable_name",
        "question_text",
        "type",
        "options",
        "example",
        "skip_if",
        "section",
    ]
    question_rows = questions or [
        ["client_name", "Client name", "text", "", "Example Ltd", "", "Matter"],
        ["urgent", "Is this urgent?", "yesno", "", "", "", "Matter"],
        ["reason", "Why?", "textarea", "", "", "urgent == false", "Matter"],
    ]
    table = document.add_table(rows=1, cols=len(question_headers))
    for cell, value in zip(table.rows[0].cells, question_headers, strict=True):
        cell.text = value
    for values in question_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value

    output_headers = ["label", "include_if", "filename_pattern"]
    output_rows = outputs or [["letter", "", "{project_code}_{label}_{date}"]]
    table = document.add_table(rows=1, cols=len(output_headers))
    for cell, value in zip(table.rows[0].cells, output_headers, strict=True):
        cell.text = value
    for values in output_rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = value
    destination = BytesIO()
    document.save(destination)
    return destination.getvalue()


def template_docx(text: str = "Dear {{ client_name }}") -> bytes:
    document = Document()
    document.add_paragraph(text)
    destination = BytesIO()
    document.save(destination)
    return destination.getvalue()

