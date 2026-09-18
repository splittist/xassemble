"""One-way export of assembly templates to highlighted, editable Word templates.

This deliberately supports the common if/elif/else subset. Unsupported control
tags fail explicitly; expressions are instructions, never evaluated as answers.
"""
from __future__ import annotations

import argparse
import re
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from jinja2 import Environment, TemplateSyntaxError, nodes
from lxml import etree

from .models import QuestionnaireSchema
from .parser import parse_questionnaire

TOKEN = re.compile(r"{{.*?}}|{%.*?%}|{#.*?#}", re.DOTALL)
COLORS = ("DDEBF7", "E2EFDA", "FCE4D6", "E4DFEC", "DDEBF0", "FBE5D6")
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


class ManualExportError(ValueError):
    """The template cannot safely be translated by this first-pass exporter."""


@dataclass
class Branch:
    number: int
    instruction: str

    @property
    def color(self) -> str:
        return COLORS[(self.number - 1) % len(COLORS)]


@dataclass
class Block:
    branch: Branch
    previous: list[int]
    has_else: bool = False


class Translator:
    def __init__(self, schema: QuestionnaireSchema):
        self.questions = {q.variable_name: q for q in schema.questions}
        self.labels = {q.variable_name: question_label(q.question_text) for q in schema.questions}
        self.labels.update({"date": "date", "project_code": "project code"})
        self.branches: list[Branch] = []
        self.stack: list[Block] = []

    @property
    def color(self) -> str | None:
        return self.stack[-1].branch.color if self.stack else None

    def expression(self, expression: str) -> nodes.Node:
        try:
            tree = Environment().parse("{{ " + expression + " }}")
        except TemplateSyntaxError as exc:
            raise ManualExportError(f"Invalid expression: {expression}") from exc
        result = tree.body[0].nodes[0]
        for name in result.find_all(nodes.Name):
            self.label(name.name)
        if isinstance(result, nodes.Name):
            self.label(result.name)
        return result

    def label(self, name: str) -> str:
        if name not in self.labels:
            raise ManualExportError(f"No questionnaire question found for variable '{name}'")
        return self.labels[name]

    def readable(self, node: nodes.Node) -> str:
        if isinstance(node, nodes.Name):
            return f'“{self.label(node.name)}”'
        if isinstance(node, nodes.Const):
            if isinstance(node.value, bool):
                return "Yes" if node.value else "No"
            if node.value is None:
                return "blank"
            return str(node.value)
        if isinstance(node, nodes.Not):
            return f"not ({self.condition(node.node)})"
        if isinstance(node, (nodes.And, nodes.Or)):
            op = "and" if isinstance(node, nodes.And) else "or"
            return f"({self.condition(node.left)}) {op} ({self.condition(node.right)})"
        if isinstance(node, nodes.Compare):
            operators = {"eq": "is", "ne": "is not", "in": "is one of", "notin": "is not one of",
                         "gt": "is greater than", "gteq": "is at least", "lt": "is less than",
                         "lteq": "is at most"}
            return self.readable(node.expr) + " " + " ".join(
                f"{operators[op.op]} {self.readable(op.expr)}" for op in node.ops
            )
        if isinstance(node, (nodes.List, nodes.Tuple)):
            return ", ".join(self.readable(item) for item in node.items)
        raise ManualExportError("This expression needs a manual instruction")

    def condition(self, node: nodes.Node) -> str:
        if isinstance(node, nodes.Name):
            question = self.questions.get(node.name)
            suffix = "is Yes" if question and question.type == "yesno" else "is provided / non-empty"
            return f"{self.readable(node)} {suffix}"
        return self.readable(node)

    def describe_condition(self, expression: str) -> str:
        node = self.expression(expression)
        try:
            return self.condition(node)
        except ManualExportError:
            return f"the following rule holds (review manually): {expression}"

    def branch(self, instruction: str) -> Branch:
        branch = Branch(len(self.branches) + 1, instruction)
        self.branches.append(branch)
        return branch

    def translate(self, token: str) -> tuple[str, bool, str | None]:
        if token.startswith("{#"):
            return "", False, self.color
        if token.startswith("{{"):
            expression = _tag_body(token)
            expression = re.sub(r"^r\s+", "", expression)
            node = self.expression(expression)
            if isinstance(node, nodes.Name):
                label = self.label(node.name)
            elif (isinstance(node, nodes.Filter) and isinstance(node.node, nodes.Name)
                  and node.name in {"upper", "lower", "title"} and not node.args
                  and not node.kwargs and node.dyn_args is None and node.dyn_kwargs is None):
                formats = {"upper": "uppercase", "lower": "lowercase", "title": "title case"}
                label = f"{self.label(node.node.name)} ({formats[node.name]})"
            else:
                label = f"Calculate / review manually: {expression}"
            return f"[{label}]", True, self.color

        command = _tag_body(token)
        command = re.sub(r"^(p|r|tr|tc)\s+", "", command)
        pieces = command.split(maxsplit=1)
        keyword = pieces[0] if pieces else ""
        expression = pieces[1] if len(pieces) > 1 else ""
        if keyword == "if":
            instruction = "Include if " + self.describe_condition(expression)
            if self.stack:
                instruction += f"; only within C{self.stack[-1].branch.number}"
            branch = self.branch(instruction)
            self.stack.append(Block(branch, [branch.number]))
            return f"[Include C{branch.number}]", False, branch.color
        if keyword in {"elif", "else"}:
            if not self.stack or self.stack[-1].has_else:
                raise ManualExportError(f"Unexpected {keyword} tag")
            block = self.stack[-1]
            prior = block.branch
            instruction = "Include only if none of " + ", ".join(f"C{i}" for i in block.previous)
            instruction += " applies"
            if keyword == "elif":
                instruction += " and " + self.describe_condition(expression)
            elif expression:
                raise ManualExportError("Unexpected expression after else")
            if len(self.stack) > 1:
                instruction += f"; only within C{self.stack[-2].branch.number}"
            branch = self.branch(instruction)
            block.branch = branch
            block.previous.append(branch.number)
            block.has_else = keyword == "else"
            return f"[End C{prior.number}] [Include C{branch.number}]", False, branch.color
        if keyword == "endif" and not expression:
            if not self.stack:
                raise ManualExportError("Unexpected endif tag")
            branch = self.stack.pop().branch
            return f"[End C{branch.number}]", False, branch.color
        raise ManualExportError(f"Unsupported manual-export tag: {token}")


def _tag_body(token: str) -> str:
    # Whitespace-control signs touch the delimiters. Do not strip arithmetic
    # signs from expressions such as {{ -5 }}.
    body = token[2:-2]
    if body.startswith(("-", "+")):
        body = body[1:]
    if body.endswith(("-", "+")):
        body = body[:-1]
    return body.strip()


def question_label(text: str) -> str:
    text = text.splitlines()[0].strip()
    return re.sub(r"^what\s+is\s+", "", text, flags=re.IGNORECASE).rstrip("? ")


def _styled_run(source: etree._Element, text: str, highlight: bool, color: str | None):
    run = etree.Element(qn("w:r"), attrib=dict(source.attrib))
    properties = source.find(qn("w:rPr"))
    properties = deepcopy(properties) if properties is not None else OxmlElement("w:rPr")
    if highlight:
        for old in properties.findall(qn("w:highlight")):
            properties.remove(old)
        element = OxmlElement("w:highlight")
        element.set(qn("w:val"), "yellow")
        properties.append(element)
    if color:
        for old in properties.findall(qn("w:shd")):
            properties.remove(old)
        element = OxmlElement("w:shd")
        element.set(qn("w:val"), "clear")
        element.set(qn("w:fill"), color)
        properties.append(element)
    if len(properties):
        run.append(properties)
    element = OxmlElement("w:t")
    element.set(qn("xml:space"), "preserve")
    element.text = text
    run.append(element)
    return run


def _paragraph(paragraph: etree._Element, translator: Translator) -> None:
    # Exclude nested text-box paragraphs; those are visited independently.
    texts = [t for t in paragraph.iter(qn("w:t"))
             if next(t.iterancestors(qn("w:p")), None) is paragraph]
    combined = "".join(t.text or "" for t in texts)
    matches = list(TOKEN.finditer(combined))
    if not matches and not translator.stack:
        if any(opener in combined for opener in ("{{", "{%", "{#")):
            raise ManualExportError("Incomplete template tag (tags must stay within a paragraph)")
        return
    # Each source character keeps its formatting. A replacement inherits the
    # first character's run, even when Word has split a tag across several runs.
    segments: list[tuple[int, int, str, bool, str | None, bool]] = []
    cursor = 0
    for match in matches:
        segments.append((cursor, match.start(), combined[cursor:match.start()],
                         False, translator.color, False))
        value, highlight, color = translator.translate(match.group())
        segments.append((match.start(), match.end(), value, highlight, color, True))
        cursor = match.end()
    segments.append((cursor, len(combined), combined[cursor:], False, translator.color, False))
    for _, _, value, _, _, replacement in segments:
        if not replacement and any(opener in value for opener in ("{{", "{%", "{#")):
            raise ManualExportError("Incomplete template tag (tags must stay within a paragraph)")
    offset = 0
    for text in texts:
        length = len(text.text or "")
        run = text.getparent()
        if run.tag != qn("w:r"):
            raise ManualExportError("Unsupported text container in template")
        replacements = []
        for start, end, value, highlight, color, replacement in segments:
            left, right = max(offset, start), min(offset + length, end)
            if left < right:
                piece = (value if offset <= start < offset + length else "") if replacement else (
                    value[left - start:right - start])
                if piece:
                    replacements.append(_styled_run(run, piece, highlight, color))
        # Split around this text child without duplicating drawings, tabs, breaks,
        # field instructions, or other non-text content in the original run.
        parent = run.getparent()
        index = parent.index(run)
        before = deepcopy(run)
        for child in list(before):
            if child.tag != qn("w:rPr"):
                before.remove(child)
        for child in list(run):
            if child is text:
                break
            if child.tag != qn("w:rPr"):
                run.remove(child)
                before.append(child)
        if any(child.tag != qn("w:rPr") for child in before):
            parent.insert(index, before)
            index += 1
        run.remove(text)
        for replacement_run in replacements:
            parent.insert(index, replacement_run)
            index += 1
        if not any(child.tag != qn("w:rPr") for child in run):
            parent.remove(run)
        offset += length


def export_manual_template(content: bytes, schema: QuestionnaireSchema, *,
                           include_if: str = "") -> bytes:
    translator = Translator(schema)
    try:
        with ZipFile(BytesIO(content)) as archive:
            parts = {item.filename: archive.read(item) for item in archive.infolist()}
            if "word/document.xml" not in parts:
                raise ManualExportError("Not a Word document")
            roots = {}
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            for name, data in parts.items():
                if not (name.startswith("word/") and name.endswith(".xml")):
                    continue
                root = etree.fromstring(data, parser)
                if not root.xpath(".//w:t", namespaces=NS):
                    continue
                roots[name] = root
                for paragraph in root.iter(qn("w:p")):
                    _paragraph(paragraph, translator)
                if translator.stack:
                    raise ManualExportError(f"Unclosed conditional block in {name}")
            body = roots.get("word/document.xml")
            if body is None:
                body = etree.fromstring(parts["word/document.xml"], parser)
                roots["word/document.xml"] = body
            _append_instructions(body, translator, include_if)
            for name, root in roots.items():
                parts[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                            standalone=True)
            destination = BytesIO()
            with ZipFile(destination, "w") as output:
                for item in archive.infolist():
                    output.writestr(item, parts[item.filename])
            return destination.getvalue()
    except (BadZipFile, etree.XMLSyntaxError, KeyError) as exc:
        raise ManualExportError("Not a readable Word template") from exc


def _append_instructions(root, translator: Translator, include_if: str) -> None:
    document = Document()
    heading = document.add_paragraph("Manual template instructions")
    heading.paragraph_format.page_break_before = True
    heading.runs[0].bold = True
    document.add_paragraph(
        "Fill in the yellow square-bracketed placeholders. Keep or delete shaded text using "
        "the rules below. Both alternatives are retained. Nested rules apply only inside their "
        "parent section; the innermost rule supplies the shading. Colours repeat: use the C "
        "numbers and Include/End markers to identify boundaries. Remove all markers, highlighting, "
        "shading and this instruction sheet when finished."
    )
    document.add_paragraph(
        "This is a manual copy. Questionnaire validation, examples and skip rules are not carried "
        "over. Review any Calculate / review manually placeholders before use."
    )
    if include_if:
        document.add_paragraph("Use this whole document only if "
                               + translator.describe_condition(include_if) + ".")
    for branch in translator.branches:
        paragraph = document.add_paragraph()
        paragraph._p.append(_styled_run(OxmlElement("w:r"),
                            f"C{branch.number}: {branch.instruction}.", False, branch.color))
    body = root.find(qn("w:body"))
    section = body.find(qn("w:sectPr"))
    for paragraph in document.paragraphs:
        if section is not None:
            section.addprevious(deepcopy(paragraph._p))
        else:
            body.append(deepcopy(paragraph._p))


def main() -> None:
    parser = argparse.ArgumentParser(description="Export a highlighted manual Word template")
    parser.add_argument("--questionnaire", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--label", help="Output label, to carry over its whole-document include rule")
    args = parser.parse_args()
    if args.output.resolve() in {args.questionnaire.resolve(), args.template.resolve()}:
        parser.error("Output must not overwrite the questionnaire or source template")
    try:
        schema = parse_questionnaire(args.questionnaire.read_bytes())
        label = args.label or args.template.stem
        output = next((item for item in schema.outputs if item.label == label), None)
        if output is None:
            parser.error("Supply --label matching an output label in the questionnaire")
        result = export_manual_template(args.template.read_bytes(), schema,
                                        include_if=output.include_if)
        args.output.write_bytes(result)
    except (ValueError, OSError) as exc:
        parser.error(str(exc))


if __name__ == "__main__":
    main()
