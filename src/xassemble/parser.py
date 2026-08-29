from __future__ import annotations

import re
from io import BytesIO

from docx import Document

from .conditions import ConditionError, referenced_variables
from .models import OutputDocument, Question, QuestionnaireSchema


class QuestionnaireError(ValueError):
    """Raised when a questionnaire document does not match the required format."""


QUESTION_HEADERS = {
    "variable_name",
    "question_text",
    "type",
    "options",
    "example",
    "skip_if",
    "section",
}
OUTPUT_HEADERS = {"label", "include_if", "filename_pattern"}
QUESTION_TYPES = {"yesno", "choice", "text", "textarea"}
IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
RESERVED_VARIABLES = {"date"}


def parse_questionnaire(content: bytes) -> QuestionnaireSchema:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise QuestionnaireError("The upload is not a readable .docx file") from exc

    question_rows: list[dict[str, str]] | None = None
    output_rows: list[dict[str, str]] | None = None
    for table in document.tables:
        headers, rows = _read_table(table)
        if QUESTION_HEADERS.issubset(headers):
            if question_rows is not None:
                raise QuestionnaireError("More than one questions table was found")
            question_rows = rows
        elif OUTPUT_HEADERS.issubset(headers):
            if output_rows is not None:
                raise QuestionnaireError("More than one output documents table was found")
            output_rows = rows

    if question_rows is None:
        raise QuestionnaireError("No table with the required question columns was found")
    if output_rows is None:
        raise QuestionnaireError("No table with the required output document columns was found")

    questions = tuple(_parse_question(row, index + 2) for index, row in enumerate(question_rows))
    outputs = tuple(_parse_output(row, index + 2) for index, row in enumerate(output_rows))
    if not questions:
        raise QuestionnaireError("The questions table must contain at least one question")
    if not outputs:
        raise QuestionnaireError("The output documents table must contain at least one output")

    variables = [question.variable_name for question in questions]
    duplicates = sorted({name for name in variables if variables.count(name) > 1})
    if duplicates:
        raise QuestionnaireError(f"Duplicate question variable(s): {', '.join(duplicates)}")
    labels = [output.label for output in outputs]
    duplicate_labels = sorted({label for label in labels if labels.count(label) > 1})
    if duplicate_labels:
        raise QuestionnaireError(f"Duplicate output label(s): {', '.join(duplicate_labels)}")

    available: set[str] = set()
    for question in questions:
        _check_references(question.skip_if, available, f"skip_if for {question.variable_name}")
        available.add(question.variable_name)
    all_variables = set(variables) | RESERVED_VARIABLES
    for output in outputs:
        _check_references(output.include_if, all_variables, f"include_if for {output.label}")
    return QuestionnaireSchema(questions=questions, outputs=outputs)


def _read_table(table: object) -> tuple[set[str], list[dict[str, str]]]:
    if not table.rows:
        return set(), []
    header_list = [_normalize(cell.text) for cell in table.rows[0].cells]
    headers = set(header_list)
    rows = []
    for table_row in table.rows[1:]:
        values = [cell.text.strip() for cell in table_row.cells]
        row = dict(zip(header_list, values, strict=False))
        if any(row.values()):
            rows.append(row)
    return headers, rows


def _normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.strip().lower()).strip("_")


def _parse_question(row: dict[str, str], row_number: int) -> Question:
    variable = row.get("variable_name", "").strip()
    raw_text = row.get("question_text", "").strip()
    text, _, commentary = raw_text.partition("\n")
    text = text.strip()
    commentary = commentary.strip()
    question_type = row.get("type", "").strip().lower()
    if not IDENTIFIER.fullmatch(variable):
        raise QuestionnaireError(f"Questions row {row_number}: invalid variable_name '{variable}'")
    if variable in RESERVED_VARIABLES:
        raise QuestionnaireError(f"Questions row {row_number}: '{variable}' is reserved")
    if not text:
        raise QuestionnaireError(f"Questions row {row_number}: question_text is required")
    if question_type not in QUESTION_TYPES:
        raise QuestionnaireError(
            f"Questions row {row_number}: type must be one of {', '.join(sorted(QUESTION_TYPES))}"
        )
    options = tuple(
        option.strip()
        for option in re.split(r"[|\n]", row.get("options", ""))
        if option.strip()
    )
    if question_type == "choice" and not options:
        raise QuestionnaireError(f"Questions row {row_number}: choice requires options")
    examples = tuple(
        example.strip()
        for example in re.split(r"[|\n]", row.get("example", ""))
        if example.strip()
    )
    return Question(
        variable_name=variable,
        question_text=text,
        type=question_type,  # type: ignore[arg-type]
        options=options,
        examples=examples,
        commentary=commentary,
        skip_if=row.get("skip_if", "").strip(),
        section=row.get("section", "").strip(),
    )


def _parse_output(row: dict[str, str], row_number: int) -> OutputDocument:
    label = row.get("label", "").strip()
    if not IDENTIFIER.fullmatch(label):
        raise QuestionnaireError(f"Outputs row {row_number}: invalid label '{label}'")
    pattern = row.get("filename_pattern", "").strip() or "{project_code}_{label}_{date}"
    return OutputDocument(
        label=label,
        include_if=row.get("include_if", "").strip(),
        filename_pattern=pattern,
    )


def _check_references(expression: str, available: set[str], location: str) -> None:
    try:
        unknown = referenced_variables(expression) - available
    except ConditionError as exc:
        raise QuestionnaireError(f"Invalid {location}: {exc}") from exc
    if unknown:
        raise QuestionnaireError(f"Unknown variable(s) in {location}: {', '.join(sorted(unknown))}")

